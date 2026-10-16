"""Deneme sonucu yükleme testleri için örnek dosyalar üretir.

Aynı sonuç tablosu, yayınevlerinin ve optik okuyucuların verdiği farklı
biçimlerde yazılır:
  sonuc.xlsx       Excel 2007+ (stiller, sayı biçimleri, birden çok sayfa)
  sonuc.xls        Excel 97-2003 (BIFF8); uzun sayfa SST'yi CONTINUE'ya böler
  sonuc_1254.csv   noktalı virgüllü, Windows-1254 (Türkçe) kodlamalı CSV
  sonuc_tab.txt    sekmeyle ayrılmış UTF-8 metin
  sonuc_html.xls   uzantısı .xls olan HTML tablo (bazı sistemlerin dışa aktarımı)
  resim.png        tablo olmayan dosya

Gerekenler: pip install xlwt==1.3.0 openpyxl==3.1.5
Kullanım:   python test/fixtures/import/make_fixtures.py
"""
import os

import openpyxl
import xlwt
from openpyxl.styles import Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))

HEADER = ['Öğrenci No', 'Ad Soyad', 'Türkçe D', 'Türkçe Y', 'Matematik D', 'Matematik Y']
ROWS = [
    ['01042', 'Elif Yılmaz', 34, 4, 26, 5],
    ['1077', 'Zeynep Arslan', 28, 6, 18, 8],
    [2001, 'Çağrı Işık', 30.0, 5, 22.5, 3],   # numara sayı, kesirli değer
    ['9999', 'Bilinmeyen Öğrenci', 20, 10, 10, 10],
]


def xlsx():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Sonuçlar'
    ws.append(['TG TYT Deneme 9 · Sonuç Listesi'])  # başlık satırı (tablodan önce)
    ws.append([])
    ws.append(HEADER)
    for c in ws[3]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='FFC531')
    for r in ROWS:
        ws.append(r)
    for row in ws.iter_rows(min_row=4, min_col=3):
        for c in row:
            c.number_format = '0.00'
    info = wb.create_sheet('Bilgi')
    info['A1'] = 'Yayınevi'
    info['B1'] = 'Test Yayınları'
    wb.save(os.path.join(HERE, 'sonuc.xlsx'))


def xls():
    wb = xlwt.Workbook(encoding='utf-8')
    ws = wb.add_sheet('Sonuçlar')
    bold = xlwt.easyxf('font: bold on')
    ws.write(0, 0, 'TG TYT Deneme 9 · Sonuç Listesi')
    for c, h in enumerate(HEADER):
        ws.write(2, c, h, bold)
    for r, row in enumerate(ROWS, start=3):
        for c, v in enumerate(row):
            ws.write(r, c, v)
    # Çok sayıda benzersiz metin: SST kaydı 8 KB'ı aşar, CONTINUE'ya bölünür.
    long = wb.add_sheet('Uzun')
    for i in range(1500):
        long.write(i, 0, f'Öğrenci {i:04d} şçğüöı')
        long.write(i, 1, i)
    wb.save(os.path.join(HERE, 'sonuc.xls'))


def csv1254():
    lines = [';'.join(HEADER)]
    for r in ROWS:
        cells = [str(v).replace('.', ',') if isinstance(v, float) else str(v) for v in r]
        lines.append(';'.join(cells))
    lines.append('5000;"Soyad, Ad";1;2;3;4')  # tırnaklı alan içinde virgül
    with open(os.path.join(HERE, 'sonuc_1254.csv'), 'wb') as f:
        f.write('\r\n'.join(lines).encode('cp1254'))


def tab_txt():
    lines = ['\t'.join(HEADER)] + ['\t'.join(str(v) for v in r) for r in ROWS]
    with open(os.path.join(HERE, 'sonuc_tab.txt'), 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))


def html_xls():
    def row(cells, tag='td'):
        return '<tr>' + ''.join(f'<{tag}>{c}</{tag}>' for c in cells) + '</tr>'
    body = row(HEADER, 'th') + ''.join(row(r) for r in ROWS)
    html = ('<html><head><meta charset="utf-8"></head><body>'
            '<table border="1">' + body.replace('Elif Yılmaz', 'Elif&nbsp;Yılmaz') +
            '</table></body></html>')
    with open(os.path.join(HERE, 'sonuc_html.xls'), 'w', encoding='utf-8') as f:
        f.write(html)


def png():
    data = bytes.fromhex(
        '89504e470d0a1a0a0000000d4948445200000001000000010806000000'
        '1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082')
    with open(os.path.join(HERE, 'resim.png'), 'wb') as f:
        f.write(data)


if __name__ == '__main__':
    xlsx()
    xls()
    csv1254()
    tab_txt()
    html_xls()
    png()
    print('ok')
